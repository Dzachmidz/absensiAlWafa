from django.core.paginator import Paginator
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import AuthenticationForm
from userapp.models import Siswa, Absensi, RfidLog, Kelas, JadwalSesi
from django.db import models
import datetime
import openpyxl
import time


def login_view(request):
    if request.method == 'POST':
        form = AuthenticationForm(request, data=request.POST)
        if form.is_valid():
            username = form.cleaned_data.get('username')
            password = form.cleaned_data.get('password')
            user = authenticate(username=username, password=password)
            if user is not None:
                login(request, user)
                messages.info(request, f"Selamat datang, {username}.")
                return redirect('dashboard')
            else:
                messages.error(request, "Username atau password salah.")
        else:
            messages.error(request, "Username atau password salah.")
    form = AuthenticationForm()
    return render(request, 'login.html', {'form': form})


def logout_view(request):
    logout(request)
    messages.info(request, "Anda berhasil logout.")
    return redirect('login')


@login_required(login_url='login')
def dashboard_view(request):
    # Hitung ringkasan data
    total_siswa = Siswa.objects.count()
    absen_hari_ini = Absensi.objects.filter(
        tanggal=datetime.date.today()).count()
    belum_hadir = total_siswa - absen_hari_ini
    if belum_hadir < 0:
        belum_hadir = 0

    # Ambil 5 aktivitas absensi terakhir hari ini (atau global jika perlu)
    recent_activity = Absensi.objects.select_related(
        'siswa', 'siswa__kelas').order_by('-tanggal', '-created_at')[:5]

    context = {
        'total_siswa': total_siswa,
        'absen_hari_ini': absen_hari_ini,
        'belum_hadir': belum_hadir,
        'recent_activity': recent_activity
    }
    return render(request, 'dashboardadmin.html', context)


@login_required(login_url='login')
def data_kelas_view(request):
    # Handle Delete
    if 'delete' in request.GET:
        try:
            id_delete = request.GET.get('delete')
            Kelas.objects.filter(id=id_delete).delete()
            messages.success(request, "Data kelas berhasil dihapus")
        except Exception as e:
            messages.error(request, f"Gagal menghapus: {e}")
        return redirect('data_kelas')

    if request.method == 'POST':
        nama_kelas = request.POST.get('nama_kelas')
        id_kelas = request.POST.get('id_kelas')  # Hidden field for edit

        # Ambil data jadwal dari form
        jmm = request.POST.get('jadwal_masuk_mulai')
        jma = request.POST.get('jadwal_masuk_akhir')
        jpm = request.POST.get('jadwal_pulang_mulai')
        jpa = request.POST.get('jadwal_pulang_akhir')

        if id_kelas:
            # === UPDATE EXISTING CLASS ===
            try:
                kelas = Kelas.objects.get(id=id_kelas)

                # Cek jika nama berubah, apakah nama baru sudah ada?
                if kelas.nama_kelas != nama_kelas and Kelas.objects.filter(nama_kelas=nama_kelas).exists():
                    messages.error(
                        request, f"Nama kelas {nama_kelas} sudah digunakan oleh kelas lain!")
                else:
                    kelas.nama_kelas = nama_kelas
                    kelas.jadwal_masuk_mulai = jmm
                    kelas.jadwal_masuk_akhir = jma
                    kelas.jadwal_pulang_mulai = jpm
                    kelas.jadwal_pulang_akhir = jpa
                    kelas.save()
                    messages.success(
                        request, f"Data kelas {nama_kelas} berhasil diperbarui")
            except Kelas.DoesNotExist:
                messages.error(request, "Kelas tidak ditemukan")

        else:
            # === CREATE NEW CLASS ===
            if nama_kelas:
                if Kelas.objects.filter(nama_kelas=nama_kelas).exists():
                    messages.error(request, f"Kelas {nama_kelas} sudah ada!")
                else:
                    Kelas.objects.create(
                        nama_kelas=nama_kelas,
                        jadwal_masuk_mulai=jmm,
                        jadwal_masuk_akhir=jma,
                        jadwal_pulang_mulai=jpm,
                        jadwal_pulang_akhir=jpa
                    )
                    messages.success(
                        request, f"Kelas {nama_kelas} berhasil ditambahkan")
        return redirect('data_kelas')

    data_kelas = Kelas.objects.prefetch_related(
        'sesi_jadwal').all().order_by('nama_kelas')
    return render(request, 'datakelas.html', {'data_kelas': data_kelas})


@login_required(login_url='login')
def kelola_sesi_view(request, kelas_id):
    if request.method == 'POST':
        action = request.POST.get('action')
        kelas = get_object_or_404(Kelas, id=kelas_id)

        try:
            if action == 'add':
                JadwalSesi.objects.create(
                    kelas=kelas,
                    nama_sesi=request.POST.get('nama_sesi'),
                    jam_masuk_mulai=request.POST.get('jam_masuk_mulai'),
                    jam_masuk_akhir=request.POST.get('jam_masuk_akhir'),
                    jam_pulang_mulai=request.POST.get('jam_pulang_mulai'),
                    jam_pulang_akhir=request.POST.get('jam_pulang_akhir')
                )
                messages.success(request, "Sesi jadwal berhasil ditambahkan")

            elif action == 'delete':
                sesi_id = request.POST.get('sesi_id')
                JadwalSesi.objects.filter(id=sesi_id, kelas=kelas).delete()
                messages.success(request, "Sesi jadwal berhasil dihapus")

        except Exception as e:
            messages.error(request, f"Error: {e}")

    return redirect('data_kelas')


@login_required(login_url='login')
def data_siswa_view(request):
    # Handle Delete
    if 'delete' in request.GET:
        try:
            id_delete = request.GET.get('delete')
            Siswa.objects.filter(id=id_delete).delete()
            messages.success(request, "Data siswa berhasil dihapus")
        except Exception as e:
            messages.error(request, f"Gagal menghapus: {e}")
        return redirect('data_siswa')

    if request.method == 'POST':
        nama = request.POST.get('nama')
        nik = request.POST.get('nik')
        kelas_id = request.POST.get('kelas_id')
        no_telepon = request.POST.get('no_telepon')

        try:
            # Validasi NIK unik
            if Siswa.objects.filter(nik=nik).exists():
                messages.error(request, f"NIK/NIS {nik} sudah terdaftar!")
            else:
                kelas_obj = Kelas.objects.get(id=kelas_id)
                Siswa.objects.create(
                    nama=nama,
                    nik=nik,
                    kelas=kelas_obj,
                    no_telepon=no_telepon
                )
                messages.success(
                    request, f"Berhasil menambahkan siswa: {nama}")
        except Exception as e:
            messages.error(request, f"Gagal menyimpan: {str(e)}")
        return redirect('data_siswa')

    data_siswa = Siswa.objects.select_related('kelas').all().order_by('-id')
    daftar_kelas = Kelas.objects.all()
    context = {
        'data_siswa': data_siswa,
        'daftar_kelas': daftar_kelas
    }
    return render(request, 'datasiswa.html', context)


@login_required(login_url='login')
def riwayat_absensi_view(request):
    # Ambil semua data absensi urutkan dari yang terbaru
    absensi_list = Absensi.objects.select_related(
        'siswa', 'siswa__kelas').all().order_by('-tanggal', '-jam_masuk')

    # Pagination handled by DataTables on client side
    return render(request, 'riwayatabsensi.html', {'data_absensi': absensi_list})


@login_required(login_url='login')
def registrasi_rfid_view(request):
    # Cek apakah ada kartu unknown terakhir (untuk auto-fill)
    last_log = RfidLog.objects.order_by('-id').first()
    suggested_uid = ""

    if last_log and last_log.status == 'unknown':
        suggested_uid = last_log.rfid_uid

    check_result = None  # Variabel untuk menampung hasil pengecekan kartu

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'check':
            # Logic Pengecekan Kartu
            rfid_check = request.POST.get('rfid_uid')
            if rfid_check:
                try:
                    siswa = Siswa.objects.get(rfid_uid=rfid_check)
                    check_result = {
                        'status': 'registered',
                        'siswa': siswa,
                        'message': 'Kartu Terdaftar'
                    }
                    messages.info(
                        request, f"Kartu terdaftar atas nama: {siswa.nama}")
                except Siswa.DoesNotExist:
                    check_result = {
                        'status': 'available',
                        'uid': rfid_check,
                        'message': 'Kartu Belum Terdaftar (Tersedia)'
                    }
                    messages.info(
                        request, "Kartu belum terdaftar. Bisa digunakan.")
        else:
            # Logic Registrasi (Existing / Baru)
            rfid_uid = request.POST.get('rfid_uid')
            siswa_id = request.POST.get('siswa_id')  # ID Siswa Existing

            # Opsi 1: Link ke siswa existing
            if siswa_id:
                try:
                    siswa = Siswa.objects.get(id=siswa_id)
                    # Cek apakah UID sudah dipakai orang lain
                    if Siswa.objects.exclude(id=siswa.id).filter(rfid_uid=rfid_uid).exists():
                        messages.error(
                            request, f"RFID {rfid_uid} sudah dipakai siswa lain!")
                    else:
                        siswa.rfid_uid = rfid_uid
                        siswa.save()
                        messages.success(
                            request, f"Kartu berhasil dihubungkan ke {siswa.nama}")
                        return redirect('registrasi_rfid')
                except Siswa.DoesNotExist:
                    messages.error(request, "Data siswa tidak ditemukan.")

            # Opsi 2: Buat siswa baru (Fallback jika admin mengisi manual data baru)
            else:
                nama = request.POST.get('nama')
                nik = request.POST.get('nik')
                kelas_id = request.POST.get('kelas_id')
                no_telepon = request.POST.get('no_telepon')

                if not nama:  # Minimal nama harus ada
                    messages.error(request, "Pilih siswa atau isi data baru!")
                elif Siswa.objects.filter(rfid_uid=rfid_uid).exists():
                    messages.error(
                        request, f"RFID UID {rfid_uid} sudah digunakan!")
                elif Siswa.objects.filter(nik=nik).exists():
                    messages.error(request, f"NIK {nik} sudah terdaftar!")
                else:
                    try:
                        cls_obj = Kelas.objects.get(
                            id=kelas_id) if kelas_id else None
                        Siswa.objects.create(
                            nama=nama,
                            nik=nik,
                            kelas=cls_obj,
                            no_telepon=no_telepon,
                            rfid_uid=rfid_uid
                        )
                        messages.success(
                            request, f"Berhasil mendaftarkan siswa baru: {nama}")
                        return redirect('registrasi_rfid')
                    except Exception as e:
                        messages.error(request, f"Gagal menyimpan: {str(e)}")

    # Ambil 5 log unknown terakhir buat referensi admin
    unknown_logs = RfidLog.objects.filter(status='unknown').order_by('-id')[:5]

    # Ambil daftar siswa yang belum punya RFID (untuk dropdown)
    siswa_tanpa_rfid = Siswa.objects.filter(
        models.Q(rfid_uid__isnull=True) | models.Q(rfid_uid='')).order_by('nama')

    # Ambil daftar kelas untuk form baru
    daftar_kelas = Kelas.objects.all()

    context = {
        'suggested_uid': suggested_uid,
        'unknown_logs': unknown_logs,
        'siswa_tanpa_rfid': siswa_tanpa_rfid,
        'daftar_kelas': daftar_kelas,
        'check_result': check_result
    }
    return render(request, 'tambahrfid.html', context)


@login_required(login_url='login')
def edit_siswa_view(request, id):
    siswa = get_object_or_404(Siswa, id=id)
    daftar_kelas = Kelas.objects.all()

    if request.method == 'POST':
        nama = request.POST.get('nama')
        nik = request.POST.get('nik')
        kelas_id = request.POST.get('kelas_id')
        no_telepon = request.POST.get('no_telepon')

        try:
            # Nik unique check (exclude self)
            if Siswa.objects.exclude(id=id).filter(nik=nik).exists():
                messages.error(
                    request, f"NIK/NIS {nik} sudah digunakan siswa lain.")
            else:
                kelas_obj = None
                if kelas_id:
                    kelas_obj = Kelas.objects.get(id=kelas_id)

                siswa.nama = nama
                siswa.nik = nik
                siswa.kelas = kelas_obj
                siswa.no_telepon = no_telepon
                siswa.save()
                messages.success(
                    request, f"Data siswa {nama} berhasil diperbarui.")
                return redirect('data_siswa')
        except Exception as e:
            messages.error(request, f"Gagal update: {e}")

    return render(request, 'edit_siswa.html', {'siswa': siswa, 'daftar_kelas': daftar_kelas})


@login_required(login_url='login')
def import_siswa_view(request):
    if request.method == 'POST':
        if 'file_excel' not in request.FILES:
            messages.error(request, "File tidak ditemukan.")
            return redirect('data_siswa')

        file_excel = request.FILES['file_excel']
        kelas_id = request.POST.get('kelas_id')

        kelas_obj = None
        if kelas_id:
            try:
                kelas_obj = Kelas.objects.get(id=kelas_id)
            except Kelas.DoesNotExist:
                messages.error(request, "Kelas tidak valid.")
                return redirect('data_siswa')

        if not file_excel.name.endswith('.xlsx'):
            messages.error(request, "Format file harus .xlsx")
            return redirect('data_siswa')

        try:
            wb = openpyxl.load_workbook(file_excel)
            ws = wb.active

            count_success = 0

            # Mencari kolom 'Nama Santri'
            header_row = next(ws.iter_rows(
                min_row=1, max_row=1, values_only=True))
            try:
                headers = [str(h).lower().strip()
                           if h else '' for h in header_row]
                nama_index = -1
                for i, h in enumerate(headers):
                    if 'nama santri' in h:
                        nama_index = i
                        break

                if nama_index == -1:
                    nama_index = 1  # Fallback B
            except:
                nama_index = 1

            current_time = int(time.time())

            for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
                if not row:
                    continue

                # Check bounds
                if len(row) <= nama_index:
                    continue

                nama = row[nama_index]
                if not nama:
                    continue

                dummy_nik = f"IMP.{current_time}.{index+1}"

                try:
                    Siswa.objects.create(
                        nama=nama,
                        nik=dummy_nik,
                        kelas=kelas_obj,
                        no_telepon="-"
                    )
                    count_success += 1
                except Exception as e:
                    print(f"Error importing row {index}: {e}")

            if count_success > 0:
                messages.success(
                    request, f"Berhasil mengimport {count_success} data siswa.")
            else:
                messages.warning(
                    request, "Tidak ada data yang berhasil diimport. Pastikan format kolom benar.")

        except Exception as e:
            messages.error(request, f"Gagal memproses file: {e}")

    return redirect('data_siswa')


@login_required(login_url='login')
def absensi_manual_view(request):
    from django.utils import timezone

    if request.method == 'POST':
        siswa_id = request.POST.get('siswa_id')
        try:
            siswa = Siswa.objects.get(id=siswa_id)

            # Logika Absensi (Diadaptasi dari RfidScanView untuk Multi-Session)
            now_aware = timezone.localtime(timezone.now())
            current_time = now_aware.time()
            today = now_aware.date()

            # Ambil Jadwal dari KELAS siswa tersebut
            kelas_siswa = siswa.kelas

            # Cek Sesi
            sessions = kelas_siswa.sesi_jadwal.all() if kelas_siswa else []
            open_absensi = Absensi.objects.filter(
                siswa=siswa, tanggal=today, jam_keluar__isnull=True).first()

            msg = ""
            status_type = "success"

            # 1. Logic PULANG (Jika ada open absensi)
            if open_absensi:
                # Tentukan jadwal pulang untuk sesi ini
                if open_absensi.sesi:
                    j_pulang_mulai = open_absensi.sesi.jam_pulang_mulai
                    j_pulang_akhir = open_absensi.sesi.jam_pulang_akhir
                    nama_sesi = open_absensi.sesi.nama_sesi
                elif kelas_siswa:
                    j_pulang_mulai = kelas_siswa.jadwal_pulang_mulai
                    j_pulang_akhir = kelas_siswa.jadwal_pulang_akhir
                    nama_sesi = "Reguler"
                else:
                    j_pulang_mulai = datetime.time(12, 0)
                    j_pulang_akhir = datetime.time(17, 0)
                    nama_sesi = ""

                # Manual override: Admin boleh memulangkan kapan saja (warning jika terlalu cepat?)
                # Tapi untuk konsistensi status, kita cek:

                status_pulang = ""
                # Cek jika terlalu cepat? (Optional)
                if current_time < j_pulang_mulai:
                    # Early departure
                    status_pulang = " (Pulang Cepat)"
                elif current_time > j_pulang_akhir:
                    status_pulang = " (Telat Jemput)"

                open_absensi.jam_keluar = current_time
                open_absensi.save()

                msg = f"Absen Pulang Berhasil: {siswa.nama} ({nama_sesi}){status_pulang}"

            # 2. Logic MASUK
            else:
                candidate_session = None

                if sessions:
                    # Cari kandidat sesi
                    valid_candidates = []
                    for ses in sessions:
                        # Untuk manual, kita lebih longgar.
                        # Jika current_time >= jam_masuk_mulai, bisa dianggap sesi ini.
                        if ses.jam_masuk_mulai <= current_time:
                            is_done = Absensi.objects.filter(
                                siswa=siswa, tanggal=today, sesi=ses).exists()
                            if not is_done:
                                valid_candidates.append(ses)

                    valid_candidates.sort(
                        key=lambda x: x.jam_masuk_mulai, reverse=True)
                    if valid_candidates:
                        candidate_session = valid_candidates[0]
                else:
                    # Default class schedule
                    if kelas_siswa:
                        is_done = Absensi.objects.filter(
                            siswa=siswa, tanggal=today).exists()
                        if not is_done:
                            candidate_session = "DEFAULT_CLASS_SCHEDULE"

                if candidate_session:
                    if candidate_session == "DEFAULT_CLASS_SCHEDULE":
                        j_masuk_akhir = kelas_siswa.jadwal_masuk_akhir
                        c_nama_sesi = ""
                        real_sesi_obj = None
                    else:
                        j_masuk_akhir = candidate_session.jam_masuk_akhir
                        c_nama_sesi = candidate_session.nama_sesi
                        real_sesi_obj = candidate_session

                    if current_time <= j_masuk_akhir:
                        status_kehadiran = "Tepat Waktu"
                    else:
                        status_kehadiran = "Terlambat"

                    Absensi.objects.create(
                        siswa=siswa,
                        tanggal=today,
                        jam_masuk=current_time,
                        hari=now_aware.strftime("%A"),
                        status_kehadiran=status_kehadiran,
                        sesi=real_sesi_obj
                    )

                    msg = f"Absen Masuk Berhasil: {siswa.nama} ({c_nama_sesi}) - {status_kehadiran}"
                else:
                    msg = "Tidak ditemukan jadwal sesi aktif atau siswa sudah absen semua sesi."
                    status_type = "warning"

            # Log Aktivitas Manual
            RfidLog.objects.create(
                rfid_uid="MANUAL-" + str(request.user.username),
                status='success',
                keterangan=f"[MANUAL] {msg} - {siswa.kelas}"
            )

            if status_type == "success":
                messages.success(request, msg)
            else:
                messages.warning(request, msg)

        except Siswa.DoesNotExist:
            messages.error(request, "Siswa tidak ditemukan.")
        except Exception as e:
            messages.error(request, f"Terjadi kesalahan: {e}")

        return redirect('absensi_manual')

    # GET Request
    siswas = Siswa.objects.select_related(
        'kelas').all().order_by('kelas__nama_kelas', 'nama')
    return render(request, 'absensi_manual.html', {'siswas': siswas})
